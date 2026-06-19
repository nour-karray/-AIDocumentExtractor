from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
ASSET_DIR = DOCS_DIR / "rapport_pfa_assets"
OUTPUT_PATH = DOCS_DIR / "RAPPORT_PFA_DETAILLE_DOCUAI.docx"

NAVY = "183B56"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "2A8C82"
GREEN = "2E8B57"
GOLD = "C28A2C"
RED = "B23A48"
INK = "1D2733"
MUTED = "5F6B76"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "EAF7F5"
LIGHT_GRAY = "F3F5F7"
LIGHT_GOLD = "FFF6DF"
WHITE = "FFFFFF"
BORDER = "C9D4DE"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[min(index, len(widths_dxa) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, size: float = 9.2) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, color=color, bold=bold)


def add_table(
    doc: Document,
    headers: list[str],
    rows: Iterable[Iterable[str]],
    widths_dxa: list[int],
    *,
    header_fill: str = NAVY,
    first_col_bold: bool = False,
) -> None:
    row_list = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, title in enumerate(headers):
        shade_cell(header.cells[index], header_fill)
        set_cell_text(header.cells[index], title, bold=True, color=WHITE, size=9.2)
        header.cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(row_list):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                shade_cell(cells[index], LIGHT_GRAY)
            set_cell_text(cells[index], value, bold=first_col_bold and index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    bold_prefix: str | None = None,
    italic: bool = False,
    align: WD_ALIGN_PARAGRAPH | None = None,
    color: str = INK,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, size=11, color=color, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(rest, size=11, color=color, italic=italic)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=color, italic=italic)


def add_bullets(doc: Document, items: Iterable[str], *, level: int = 0) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.18
        run = paragraph.add_run(item)
        set_run_font(run, size=10.7, color=INK)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.18
        run = paragraph.add_run(item)
        set_run_font(run, size=10.7, color=INK)


def add_callout(doc: Document, title: str, text: str, *, fill: str = LIGHT_BLUE, accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    title_run = paragraph.add_run(title)
    set_run_font(title_run, size=10.5, color=accent, bold=True)
    body = cell.add_paragraph()
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.line_spacing = 1.18
    body_run = body.add_run(text)
    set_run_font(body_run, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, "F6F8FA")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_run_font(run, name="Consolas", size=8.6, color="263238")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_chapter(doc: Document, number: int, title: str) -> None:
    doc.add_page_break()
    add_heading(doc, f"Chapitre {number} - {title}", 1)


def add_field(paragraph, instruction: str) -> None:
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)

    instruction_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instruction_run.append(instr)

    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)

    value_run = OxmlElement("w:r")
    value_text = OxmlElement("w:t")
    value_text.text = "1"
    value_run.append(value_text)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)

    paragraph._p.append(begin_run)
    paragraph._p.append(instruction_run)
    paragraph._p.append(separate_run)
    paragraph._p.append(value_run)
    paragraph._p.append(end_run)


def set_picture_alt(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def font(size: int, bold: bool = False):
    regular = Path(r"C:\Windows\Fonts\arial.ttf")
    bold_path = Path(r"C:\Windows\Fonts\arialbd.ttf")
    path = bold_path if bold and bold_path.is_file() else regular
    if path.is_file():
        return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline, radius=24, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered_text(draw, box, text, font_obj, fill, spacing=6):
    left, top, right, bottom = box
    wrapped: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        line = ""
        for word in words:
            candidate = f"{line} {word}".strip()
            if draw.textlength(candidate, font=font_obj) <= (right - left - 30):
                line = candidate
            else:
                if line:
                    wrapped.append(line)
                line = word
        if line:
            wrapped.append(line)
    line_height = font_obj.getbbox("Ag")[3] - font_obj.getbbox("Ag")[1]
    total = len(wrapped) * line_height + max(0, len(wrapped) - 1) * spacing
    y = top + ((bottom - top - total) / 2)
    for line in wrapped:
        width = draw.textlength(line, font=font_obj)
        draw.text((left + ((right - left - width) / 2), y), line, font=font_obj, fill=fill)
        y += line_height + spacing


def draw_arrow(draw, start, end, color=f"#{NAVY}", width=5):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    if abs(end[0] - start[0]) >= abs(end[1] - start[1]):
        direction = 1 if end[0] > start[0] else -1
        points = [(x2, y2), (x2 - 16 * direction, y2 - 10), (x2 - 16 * direction, y2 + 10)]
    else:
        direction = 1 if end[1] > start[1] else -1
        points = [(x2, y2), (x2 - 10, y2 - 16 * direction), (x2 + 10, y2 - 16 * direction)]
    draw.polygon(points, fill=color)


def build_architecture_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 930), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(42, True)
    box_title = font(25, True)
    box_body = font(20)
    draw.text((70, 35), "Architecture fonctionnelle de DocuAI", font=title_font, fill=f"#{NAVY}")

    boxes = {
        "ui": (70, 135, 390, 330),
        "api": (510, 135, 830, 330),
        "orch": (950, 135, 1530, 330),
        "ocr": (70, 510, 400, 760),
        "local": (490, 510, 900, 760),
        "gemini": (990, 510, 1320, 760),
        "store": (1370, 510, 1530, 760),
    }
    styles = {
        "ui": ("#EAF2F8", "#2E74B5"),
        "api": ("#EAF7F5", "#2A8C82"),
        "orch": ("#FFF6DF", "#C28A2C"),
        "ocr": ("#F3F5F7", "#5F6B76"),
        "local": ("#EEEAFB", "#7057B7"),
        "gemini": ("#E8F4FF", "#3F7CC4"),
        "store": ("#EAF7F0", "#2E8B57"),
    }
    labels = {
        "ui": ("Frontend Next.js / React", "Upload, configuration,\napercu, resultats,\nhistorique, evaluation"),
        "api": ("API FastAPI", "Endpoints REST, auth,\nCORS, validation,\ntelechargements"),
        "orch": ("Orchestrateur Python", "Detection du type,\nselection du pipeline,\ntimeouts, normalisation"),
        "ocr": ("OCR local specialise", "OpenCV + Tesseract\n+ regles medicales,\nSTEG, tickets"),
        "local": ("Pipeline IA local", "Pretraitement -> Docling\n-> controle qualite\n-> PaddleOCR -> Qwen2.5"),
        "gemini": ("Pipeline cloud", "Gemini Vision\nvers JSON structure\navec retries"),
        "store": ("Persistance", "SQLite\nJSON\nBLOB\nPDF/ZIP"),
    }
    for key, box in boxes.items():
        fill, outline = styles[key]
        rounded_box(draw, box, fill, outline)
        title, body = labels[key]
        left, top, right, bottom = box
        centered_text(draw, (left + 10, top + 12, right - 10, top + 75), title, box_title, outline)
        centered_text(draw, (left + 12, top + 80, right - 12, bottom - 10), body, box_body, "#263238")

    draw_arrow(draw, (390, 232), (510, 232))
    draw_arrow(draw, (830, 232), (950, 232))
    draw_arrow(draw, (1120, 330), (250, 510))
    draw_arrow(draw, (1180, 330), (695, 510))
    draw_arrow(draw, (1240, 330), (1155, 510))
    draw_arrow(draw, (1320, 635), (1370, 635))
    draw_arrow(draw, (900, 635), (1370, 635))
    draw_arrow(draw, (400, 635), (1370, 700))

    draw.text((70, 840), "Flux principal : document -> extraction -> JSON valide -> historique exploitable",
              font=font(24, True), fill=f"#{DARK_BLUE}")
    image.save(path)


def build_pipeline_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1050), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 35), "Pipeline IA local hybride", font=font(42, True), fill=f"#{NAVY}")

    steps = [
        ("1", "Validation", "Format, taille,\nfichier non vide"),
        ("2", "Pretraitement", "Orientation, perspective,\ndeskew, contraste"),
        ("3", "Docling", "Conversion vers\nMarkdown structure"),
        ("4", "Controle qualite", "Longueur, chiffres,\nstructure, lisibilite"),
        ("5", "PaddleOCR", "Fallback si le\ncontenu est faible"),
        ("6", "Qwen2.5", "Generation JSON\nvia Ollama"),
        ("7", "Validation metier", "Champs obligatoires,\ncoherence, warnings"),
        ("8", "Archivage", "SQLite + JSON +\nsource + rapport"),
    ]
    colors = ["#EAF2F8", "#EAF7F5", "#EEEAFB", "#FFF6DF", "#FCEFEF", "#E8F4FF", "#EAF7F0", "#F3F5F7"]
    outlines = ["#2E74B5", "#2A8C82", "#7057B7", "#C28A2C", "#B23A48", "#3F7CC4", "#2E8B57", "#5F6B76"]
    positions = [
        (70, 170, 390, 360),
        (460, 170, 780, 360),
        (850, 170, 1170, 360),
        (1240, 170, 1530, 360),
        (1240, 600, 1530, 790),
        (850, 600, 1170, 790),
        (460, 600, 780, 790),
        (70, 600, 390, 790),
    ]
    for index, ((number, title, body), box) in enumerate(zip(steps, positions)):
        rounded_box(draw, box, colors[index], outlines[index])
        left, top, right, bottom = box
        draw.ellipse((left + 16, top + 16, left + 62, top + 62), fill=outlines[index])
        number_width = draw.textlength(number, font=font(22, True))
        draw.text((left + 39 - number_width / 2, top + 22), number, font=font(22, True), fill="white")
        centered_text(draw, (left + 65, top + 15, right - 10, top + 72), title, font(24, True), outlines[index])
        centered_text(draw, (left + 15, top + 82, right - 15, bottom - 10), body, font(20), "#263238")

    draw_arrow(draw, (390, 265), (460, 265))
    draw_arrow(draw, (780, 265), (850, 265))
    draw_arrow(draw, (1170, 265), (1240, 265))
    draw_arrow(draw, (1385, 360), (1385, 600))
    draw_arrow(draw, (1240, 695), (1170, 695))
    draw_arrow(draw, (850, 695), (780, 695))
    draw_arrow(draw, (460, 695), (390, 695))
    draw.text((1070, 445), "Si Markdown insuffisant", font=font(22, True), fill=f"#{RED}")
    draw.text((625, 445), "Sinon passage direct vers Qwen", font=font(22, True), fill=f"#{GREEN}")
    image.save(path)


def build_dataset_chart(path: Path) -> None:
    image = Image.new("RGB", (1500, 800), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 40), "Repartition du dataset prepare", font=font(42, True), fill=f"#{NAVY}")
    data = [("Train", 839, "#2E74B5"), ("Validation", 178, "#C28A2C"), ("Test", 182, "#2E8B57")]
    max_value = max(value for _, value, _ in data)
    y_positions = [210, 390, 570]
    for (label, value, color), y in zip(data, y_positions):
        draw.text((80, y), label, font=font(28, True), fill="#263238")
        draw.rounded_rectangle((310, y, 1320, y + 70), radius=18, fill="#EEF1F4")
        width = int(1010 * (value / max_value))
        draw.rounded_rectangle((310, y, 310 + width, y + 70), radius=18, fill=color)
        draw.text((1345, y + 12), str(value), font=font(30, True), fill=color)
    draw.text((80, 720), "Total : 1 199 exemples | Medical : 226 | Tickets : 973",
              font=font(26, True), fill=f"#{DARK_BLUE}")
    image.save(path)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.28

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (11.5, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[list_style]
        style.font.name = "Calibri"
        style.font.size = Pt(10.7)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.18

    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("DocuAI - Rapport PFA detaille")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Projet de fin d'etudes 2025-2026  |  Page ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(p, "PAGE")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PROJET DE FIN D'ETUDES")
    set_run_font(run, size=13, color=TEAL, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(38)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Conception et developpement d'une plateforme intelligente")
    set_run_font(run, size=24, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("d'extraction d'informations a partir de documents heterogenes")
    set_run_font(run, size=21, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DocuAI / pfaEXTRACT")
    set_run_font(run, size=28, color=TEAL, bold=True)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2600, 6760])
    set_repeat_table_header(table.rows[0])
    labels = [
        ("Realise par", "[Nom et prenom]"),
        ("Encadre par", "[Nom de l'encadrant]"),
        ("Etablissement", "[Nom de l'etablissement]"),
        ("Annee universitaire", "2025-2026"),
    ]
    for row, (label, value) in zip(table.rows, labels):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, bold=True, color=NAVY, size=10.5)
        set_cell_text(row.cells[1], value, size=10.5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Rapport etabli a partir de l'etat du depot au 15 juin 2026")
    set_run_font(run, size=10, color=MUTED, italic=True)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    add_heading(doc, "Resume", 1)
    add_paragraph(
        doc,
        "Ce projet porte sur la conception et le developpement d'une plateforme intelligente "
        "capable de transformer des documents heterogenes en donnees JSON structurees et "
        "exploitables. La solution traite des images et des fichiers PDF appartenant a quatre "
        "familles principales : analyses medicales, factures STEG, tickets de caisse et factures "
        "fournisseurs. L'architecture actuelle associe une interface Next.js/React, une API "
        "FastAPI, des services Python de pretraitement et de routage, trois strategies "
        "d'extraction complementaires et une persistance SQLite/JSON."
    )
    add_paragraph(
        doc,
        "La strategie retenue est hybride. Un OCR local specialise reste adapte aux documents "
        "relativement stables et explicables. Un pipeline IA local combine le pretraitement, "
        "Docling, un controle de qualite, PaddleOCR en secours et Qwen2.5 via Ollama. Gemini "
        "constitue une alternative cloud pour les mises en page plus variables. Les resultats "
        "sont normalises, valides, notes par completude, archives avec le document source et "
        "restitues dans une interface comprenant tableau de bord, historique, details, exports "
        "PDF/ZIP et page d'evaluation."
    )
    add_paragraph(
        doc,
        "Le projet comprend egalement une chaine de preparation de donnees pour un futur "
        "fine-tuning LoRA de Qwen2.5. Le dataset prepare contient 1 199 exemples, repartis en "
        "839 exemples d'entrainement, 178 de validation et 182 de test. Un smoke test controle "
        "a valide la detection de quatre familles sur quatre cas, tandis que les metriques "
        "globales Precision, Recall et F1 restent volontairement non annoncees tant que les "
        "predictions comparables des pipelines n'ont pas ete produites sur le split de test."
    )
    add_paragraph(doc, "Mots-cles : OCR, extraction d'information, Document AI, Next.js, React, FastAPI, Docling, PaddleOCR, Qwen2.5, Ollama, Gemini, SQLite, LoRA.", italic=True)

    add_heading(doc, "Abstract", 1)
    add_paragraph(
        doc,
        "This project designs and implements an intelligent platform that converts heterogeneous "
        "documents into structured JSON data. The system supports medical laboratory reports, "
        "STEG utility bills, receipts and supplier invoices. Its current architecture combines "
        "a Next.js/React frontend, a FastAPI backend, Python preprocessing and routing services, "
        "three complementary extraction strategies, and SQLite/JSON persistence. The solution "
        "also prepares a supervised dataset for a future Qwen2.5 LoRA fine-tuning workflow. "
        "Measured functional results are reported separately from metrics that are still pending "
        "in order to preserve experimental integrity."
    )

    add_heading(doc, "Sommaire", 1)
    chapters = [
        "Introduction generale",
        "Contexte, problematique et objectifs",
        "Etat de l'art et choix technologiques",
        "Methodologie et organisation du travail",
        "Analyse des besoins et cas d'utilisation",
        "Conception et architecture du systeme",
        "Realisation technique",
        "Preparation des donnees et fine-tuning",
        "Tests, evaluation et resultats",
        "Limites, securite et perspectives",
        "Conclusion generale",
        "Bibliographie et annexes",
    ]
    add_numbered(doc, chapters)

    add_heading(doc, "Liste des abreviations", 1)
    add_table(
        doc,
        ["Abreviation", "Signification"],
        [
            ("API", "Application Programming Interface"),
            ("OCR", "Optical Character Recognition"),
            ("LLM", "Large Language Model"),
            ("VLM", "Vision-Language Model"),
            ("JSON", "JavaScript Object Notation"),
            ("JWT", "JSON Web Token"),
            ("SFT", "Supervised Fine-Tuning"),
            ("LoRA", "Low-Rank Adaptation"),
            ("CER", "Character Error Rate"),
            ("WER", "Word Error Rate"),
            ("PFA/PFE", "Projet de fin d'annee / Projet de fin d'etudes"),
        ],
        [2100, 7260],
        first_col_bold=True,
    )


def build_report() -> Path:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    architecture_path = ASSET_DIR / "architecture_docuai.png"
    pipeline_path = ASSET_DIR / "pipeline_local.png"
    dataset_path = ASSET_DIR / "dataset_splits.png"
    build_architecture_diagram(architecture_path)
    build_pipeline_diagram(pipeline_path)
    build_dataset_chart(dataset_path)

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc)

    add_chapter(doc, 1, "Introduction generale")
    add_paragraph(
        doc,
        "La transformation numerique produit un volume croissant de documents non structures : "
        "factures, tickets, analyses medicales, formulaires et pieces scannees. Ces documents "
        "contiennent des informations utiles, mais leur exploitation manuelle est lente, "
        "couteuse et sujette aux erreurs. Le projet DocuAI repond a ce probleme en proposant une "
        "chaine complete allant de l'import du document jusqu'a l'affichage, l'historisation et "
        "l'export des donnees extraites."
    )
    add_paragraph(
        doc,
        "Le principal enjeu n'est pas seulement de lire du texte. Il faut reconnaitre le type du "
        "document, conserver sa structure, extraire les champs metier attendus, eviter "
        "l'invention de valeurs, signaler les incertitudes et rendre le resultat exploitable par "
        "une application. Pour cette raison, le projet compare et combine plusieurs pipelines "
        "plutot que de dependre d'un seul modele."
    )
    add_callout(
        doc,
        "Question centrale",
        "Comment automatiser l'extraction d'informations fiables depuis des images et PDF "
        "heterogenes, tout en conciliant precision, cout, confidentialite, explicabilite et "
        "simplicite d'utilisation ?",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )
    add_heading(doc, "1.1 Contributions principales", 2)
    add_bullets(
        doc,
        [
            "Une application web moderne en Next.js et React pour importer, configurer, visualiser et administrer les extractions.",
            "Une API FastAPI reutilisant les services Python historiques du projet.",
            "Un routeur heuristique capable d'orienter un document vers la famille STEG, medicale, ticket ou facture fournisseur.",
            "Un pretraitement d'image comprenant validation, correction d'orientation, rectification de perspective, deskew et amelioration visuelle.",
            "Trois strategies d'extraction : OCR local specialise, IA locale hybride et Gemini API.",
            "Des schemas JSON normalises, une validation metier, des warnings et un score de completude par famille.",
            "Une persistance double SQLite/JSON, incluant le document source, la corbeille et la generation de rapports PDF/ZIP.",
            "Un dataset train/validation/test et des scripts pour l'evaluation et un futur fine-tuning LoRA de Qwen2.5.",
        ],
    )

    add_chapter(doc, 2, "Contexte, problematique et objectifs")
    add_heading(doc, "2.1 Contexte documentaire", 2)
    add_paragraph(
        doc,
        "Les documents cibles varient fortement. Une facture STEG possede des zones relativement "
        "stables, alors qu'un ticket de caisse change selon l'enseigne, le pays et la qualite de "
        "l'impression. Une analyse medicale contient des tableaux, des unites et des intervalles "
        "de reference. Une facture fournisseur peut etre multilingue et presenter des lignes, "
        "taxes, remises et parties commerciales complexes. Cette diversite justifie une "
        "architecture modulaire et un routage adapte."
    )
    add_heading(doc, "2.2 Problemes identifies", 2)
    add_bullets(
        doc,
        [
            "Qualite variable des photos : perspective, rotation, ombres, flou, faible contraste.",
            "Documents PDF natifs ou scans ne proposant pas la meme qualite de texte.",
            "Bruit OCR, notamment sur les montants, dates, numeros et textes arabes.",
            "Mises en page heterogenes qui limitent les expressions regulieres fixes.",
            "Risque d'hallucination des modeles generatifs.",
            "Dependance possible a une API externe, avec cout, quota, latence et confidentialite.",
            "Temps de calcul important pour un modele local sur une machine sans GPU adapte.",
            "Necessite d'une trace technique pour expliquer les fallbacks et les erreurs.",
        ],
    )
    add_heading(doc, "2.3 Objectif general", 2)
    add_paragraph(
        doc,
        "L'objectif general est de construire une plateforme de bout en bout capable de recevoir "
        "un ou plusieurs documents, d'identifier leur nature, d'executer la meilleure strategie "
        "disponible, de produire un JSON metier valide et de conserver le resultat dans un "
        "historique consultable."
    )
    add_heading(doc, "2.4 Objectifs specifiques", 2)
    add_numbered(
        doc,
        [
            "Supporter les formats PDF, JPG, JPEG, PNG, TIFF, WEBP et BMP jusqu'a 20 Mo.",
            "Offrir un mode automatique et des modes forces par famille documentaire.",
            "Permettre le choix entre moteur local, OCR classique et Gemini.",
            "Produire des champs adaptes a chaque type de document.",
            "Enregistrer les succes et les echecs pour faciliter l'audit.",
            "Afficher des indicateurs d'activite, de statut et de completude.",
            "Preparer une evaluation reproductible sur une verite terrain gelee.",
            "Preparer les donnees necessaires a l'adaptation future de Qwen2.5.",
        ],
    )

    add_chapter(doc, 3, "Etat de l'art et choix technologiques")
    add_heading(doc, "3.1 OCR classique", 2)
    add_paragraph(
        doc,
        "L'OCR transforme une image en texte. Tesseract est pertinent pour une execution locale, "
        "gratuite et explicable. Il devient particulierement efficace lorsque le document possede "
        "une structure connue et que le pretraitement est adapte. En revanche, l'OCR brut ne "
        "comprend pas toujours la relation semantique entre une etiquette et sa valeur."
    )
    add_heading(doc, "3.2 Vision par ordinateur", 2)
    add_paragraph(
        doc,
        "OpenCV est utilise avant l'OCR pour ameliorer la qualite d'entree. Les techniques "
        "employees comprennent la detection de contours, la recherche d'un quadrilatere, la "
        "transformation de perspective, l'analyse des lignes pour le deskew, le redimensionnement, "
        "le contraste et la nettete. Le pretraitement diminue le bruit transmis aux moteurs "
        "suivants."
    )
    add_heading(doc, "3.3 Conversion documentaire structuree", 2)
    add_paragraph(
        doc,
        "Docling joue le role de convertisseur documentaire. Il cherche a conserver l'ordre de "
        "lecture, les blocs, les titres et les tableaux dans une representation Markdown. Il ne "
        "doit pas etre presente comme un simple modele OCR : sa valeur est la production d'un "
        "contenu structure plus facile a fournir a un LLM."
    )
    add_heading(doc, "3.4 Modeles de langage et Vision-Language Models", 2)
    add_paragraph(
        doc,
        "Qwen2.5 est execute localement via Ollama. Le backend lui transmet le contenu du "
        "document, le type cible et un schema JSON strict. La temperature est fixee a zero afin "
        "de favoriser la stabilite. Gemini Vision constitue une reference cloud capable de lire "
        "directement les images et de produire une structure metier sur des documents complexes."
    )
    add_heading(doc, "3.5 Justification des choix", 2)
    add_table(
        doc,
        ["Technologie", "Role", "Justification"],
        [
            ("Next.js 15 / React 19", "Interface utilisateur", "Composants reutilisables, routage App Router, experience web moderne."),
            ("TypeScript", "Fiabilite frontend", "Contrats de types pour les payloads API et reduction des erreurs."),
            ("Tailwind CSS", "Design", "Mise en page responsive, themes clair/sombre et composants coherents."),
            ("FastAPI", "API REST", "Validation, documentation automatique, upload multipart et integration Python."),
            ("Pydantic", "Schemas", "Normalisation et validation des structures medicales, tickets, STEG et fournisseurs."),
            ("OpenCV", "Pretraitement", "Correction geometrique et amelioration d'image."),
            ("Tesseract / PaddleOCR", "Lecture de texte", "Baseline locale et fallback multilingue."),
            ("Docling", "Structure documentaire", "Conversion de PDF/images en contenu structure."),
            ("Qwen2.5 / Ollama", "Extraction locale", "Confidentialite et absence de cout API par appel."),
            ("Gemini 2.5 Flash", "Extraction cloud", "Compréhension visuelle des documents variables."),
            ("SQLite + JSON", "Persistance", "Simplicite locale, audit lisible et portabilite."),
        ],
        [1700, 2300, 5360],
        first_col_bold=True,
    )

    add_chapter(doc, 4, "Methodologie et organisation du travail")
    add_heading(doc, "4.1 Approche iterative", 2)
    add_paragraph(
        doc,
        "Le projet a ete construit par iterations courtes. Chaque iteration ajoute une capacite "
        "observable : extraction d'une famille, historisation, interface, API, pipeline local, "
        "authentification, evaluation puis preparation du fine-tuning. Cette organisation se "
        "rapproche de Scrum pour la planification incrementale et de pratiques XP pour le "
        "refactoring, les tests frequents et la livraison continue d'une version executable."
    )
    add_heading(doc, "4.2 Demarche Data Science inspiree de CRISP-DM", 2)
    add_numbered(
        doc,
        [
            "Compréhension métier : identification des champs utiles pour chaque famille.",
            "Compréhension des données : observation des images, PDF, annotations et erreurs OCR.",
            "Préparation : nettoyage, normalisation, conversion des documents et création des splits.",
            "Modélisation : règles OCR, prompts structurés, choix du pipeline et schémas Pydantic.",
            "Évaluation : smoke tests, QA fonctionnelle, métriques champ par champ et analyse d'erreurs.",
            "Déploiement applicatif : exposition FastAPI, interface Next.js, historique et exports.",
        ],
    )
    add_heading(doc, "4.3 Chronologie technique du travail", 2)
    add_table(
        doc,
        ["Phase", "Travail realise", "Livrable"],
        [
            ("1. Exploration", "Collecte de documents et definition des champs.", "Jeux d'images et schemas cibles."),
            ("2. Prototype OCR", "Extraction medicale et STEG avec Tesseract/OpenCV.", "Services Python specialises."),
            ("3. Integration Gemini", "Prompts Vision et sorties JSON.", "Pipelines cloud par famille."),
            ("4. Historique", "Sauvegarde JSON puis SQLite, rapports PDF.", "Audit et consultation."),
            ("5. Application web", "Migration vers FastAPI + Next.js/React.", "Interface admin multi-pages."),
            ("6. IA locale", "Docling, PaddleOCR, Qwen2.5 et Ollama.", "Pipeline hybride local."),
            ("7. Robustesse", "Timeouts, fallbacks, warnings, corbeille, auth.", "Application stabilisee."),
            ("8. Evaluation", "QA, splits, métriques et scripts LoRA.", "Protocole experimental reproductible."),
        ],
        [1450, 5000, 2910],
        first_col_bold=True,
    )
    add_heading(doc, "4.4 Principe d'integrite experimentale", 2)
    add_callout(
        doc,
        "Regle de communication",
        "Un score de completude affiche dans l'application n'est pas une accuracy scientifique. "
        "Les metriques Precision, Recall et F1 ne sont annoncees que lorsqu'une prediction brute "
        "est comparee a une verite terrain gelee.",
        fill=LIGHT_GOLD,
        accent=GOLD,
    )

    add_chapter(doc, 5, "Analyse des besoins et cas d'utilisation")
    add_heading(doc, "5.1 Acteurs", 2)
    add_bullets(
        doc,
        [
            "Administrateur ou operateur : importe les documents, choisit une methode et consulte les resultats.",
            "Backend DocuAI : valide, route, traite, normalise et archive.",
            "Ollama/Qwen local : produit un JSON a partir du contenu structure.",
            "Gemini API : execute l'extraction cloud lorsque ce mode est selectionne.",
            "Base SQLite et systeme de fichiers : conservent les extractions et les sources.",
        ],
    )
    add_heading(doc, "5.2 Besoins fonctionnels", 2)
    add_table(
        doc,
        ["ID", "Besoin", "Implementation"],
        [
            ("BF01", "S'authentifier ou creer un compte.", "Login/register, token Bearer optionnel."),
            ("BF02", "Importer un ou plusieurs documents.", "Upload multipart et glisser-deposer."),
            ("BF03", "Choisir auto, medical, STEG, fournisseur ou ticket.", "Modes exposes par /api/meta."),
            ("BF04", "Choisir local, OCR ou Gemini.", "Parametre method du formulaire."),
            ("BF05", "Visualiser le document et les champs.", "Apercu image/PDF et composant ResultDetail."),
            ("BF06", "Rechercher et filtrer l'historique.", "Filtres type, texte, date et pagination."),
            ("BF07", "Exporter un rapport ou plusieurs rapports.", "PDF individuel et ZIP multi-selection."),
            ("BF08", "Supprimer sans perte immediate.", "Deplacement vers Data/history/trash."),
            ("BF09", "Consulter les modeles et leur disponibilite.", "Endpoint /api/models et detection runtime."),
            ("BF10", "Consulter le protocole d'evaluation.", "Page Evaluation IA et scripts de mesures."),
        ],
        [950, 3650, 4760],
    )
    add_heading(doc, "5.3 Besoins non fonctionnels", 2)
    add_bullets(
        doc,
        [
            "Performance : limiter les blocages par des timeouts et des branches rapides.",
            "Robustesse : archiver aussi les echecs et produire des messages comprehensibles.",
            "Securite : ne pas committer les secrets, hacher les mots de passe et verifier les tokens.",
            "Confidentialite : proposer une execution locale sans transfert cloud.",
            "Maintenabilite : separer frontend, API, services, extraction, scripts et donnees.",
            "Traçabilite : conserver la methode, le type detecte, les warnings et la source.",
            "Ergonomie : interface responsive, themes et apercu direct des documents.",
        ],
    )

    add_chapter(doc, 6, "Conception et architecture du systeme")
    add_heading(doc, "6.1 Architecture generale", 2)
    architecture_shape = doc.add_picture(str(architecture_path), width=Cm(16.2))
    set_picture_alt(
        architecture_shape,
        "Architecture fonctionnelle de DocuAI",
        "Schema reliant le frontend Next.js, l'API FastAPI, l'orchestrateur, les pipelines OCR, local et Gemini, puis la persistance.",
    )
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Figure 1 - Architecture fonctionnelle de DocuAI")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.runs[0], size=9, color=MUTED, italic=True)
    add_paragraph(
        doc,
        "L'architecture est organisee en couches. Le frontend ne connait pas les details des "
        "moteurs d'extraction : il consomme des contrats TypeScript correspondant aux payloads "
        "FastAPI. Le backend orchestre les services existants, tandis que les modules Python "
        "encapsulent le routage, le pretraitement, l'OCR, les appels IA, la validation et le stockage."
    )
    add_heading(doc, "6.2 Flux principal", 2)
    add_numbered(
        doc,
        [
            "Le navigateur construit un FormData contenant les fichiers et la configuration.",
            "FastAPI lit les fichiers et transfere le traitement lourd dans un threadpool.",
            "Le document est copie dans un fichier temporaire et son type est detecte.",
            "Le pipeline choisi produit un payload JSON normalise.",
            "Le backend calcule warnings, resume et score de completude.",
            "Le resultat et la source sont enregistres dans SQLite et dans un JSON lisible.",
            "L'API retourne une cle d'historique et les URLs de detail, source et rapport.",
            "Le frontend redirige vers la page Documents et affiche le resultat.",
        ],
    )
    add_heading(doc, "6.3 Detection automatique", 2)
    add_paragraph(
        doc,
        "Le routeur exploite d'abord le nom d'origine et le chemin logique. Cette precaution est "
        "importante car le backend travaille sur un nom temporaire aleatoire. Si aucun indice "
        "fiable n'est disponible, un texte leger est produit et plusieurs scores ponderes sont "
        "calcules. Les mots STEG, compteur et kWh favorisent la facture d'electricite; les termes "
        "laboratoire, hemoglobine ou TSH favorisent le medical; les mots caisse, paiement et "
        "merci favorisent le ticket; invoice, supplier, TVA, IBAN ou leurs equivalents arabes "
        "favorisent la facture fournisseur."
    )
    add_heading(doc, "6.4 Schemas metier", 2)
    add_table(
        doc,
        ["Famille", "Champs principaux"],
        [
            ("Analyse medicale", "Laboratoire, medecin, patient, dossier, dates, tests, valeurs, unites, references, statut."),
            ("Facture STEG", "Reference, compteur, date facture, montant, echeance, periode et coupon."),
            ("Ticket de caisse", "Magasin, date, heure, numero, devise, articles, total et paiement."),
            ("Facture fournisseur", "Numero, dates, devise, vendeur, client, articles, taxes, totaux et montant du."),
        ],
        [2200, 7160],
        first_col_bold=True,
    )
    add_heading(doc, "6.5 Persistance", 2)
    add_paragraph(
        doc,
        "La table extraction_history contient l'identifiant, le type, le nom source, la date, le "
        "chemin relatif, le JSON, le statut, le message d'erreur, le type detecte, le MIME, la "
        "taille et le document source en BLOB. Un fichier JSON et une copie du document sont "
        "egalement ecrits dans un sous-dossier par type. Cette duplication facilite le debug et "
        "l'export, mais doit etre surveillee lorsque le volume augmente."
    )

    add_chapter(doc, 7, "Realisation technique")
    add_heading(doc, "7.1 Frontend Next.js et React", 2)
    add_paragraph(
        doc,
        "Le frontend utilise l'App Router de Next.js avec des pages TypeScript. Les pages "
        "interactives sont des composants client et utilisent useState pour les formulaires, "
        "useEffect pour charger les donnees et useMemo pour les calculs derives. L'interface "
        "partage un AppShell responsable de la navigation, du theme et de la deconnexion."
    )
    add_table(
        doc,
        ["Page", "Responsabilite"],
        [
            ("/login et /register", "Authentification, stockage du token et redirection."),
            ("/dashboard", "Activite recente, graphiques et score qualite moyen."),
            ("/documents", "Galerie, selection, detail, export PDF/ZIP et corbeille."),
            ("/extractions", "Upload, apercu, choix du type et du moteur, lancement du batch."),
            ("/history", "Historique groupe par date avec recherche et filtres."),
            ("/results", "Affichage du dernier resultat ou d'une entree cible."),
            ("/evaluation", "Dataset, metriques, pipelines et limites experimentales."),
            ("/settings", "Theme, fournisseur IA, modele et configuration locale navigateur."),
        ],
        [2300, 7060],
        first_col_bold=True,
    )
    add_heading(doc, "7.2 Communication avec l'API", 2)
    add_paragraph(
        doc,
        "Le module frontend/lib/api.ts centralise les appels fetch, l'ajout de l'en-tete Bearer, "
        "la lecture des erreurs FastAPI et les telechargements binaires. L'upload utilise "
        "AbortController avec une limite navigateur de 150 secondes. Les preferences de theme, "
        "methode et modele sont conservees dans localStorage; le token est conserve dans "
        "sessionStorage."
    )
    add_code_block(
        doc,
        "POST /api/extractions\n"
        "FormData: files[], mode, method, geminiApiKey, geminiModel,\n"
        "          ollamaHost, localModel, retries, retryDelay, originsJson",
    )
    add_heading(doc, "7.3 Backend FastAPI", 2)
    add_paragraph(
        doc,
        "FastAPI expose les endpoints de sante, metadata, authentification, dashboard, modeles, "
        "historique, source, rapports et extraction. Les traitements synchrones lourds sont "
        "executes hors de la boucle asynchrone avec run_in_threadpool. Chaque document possede "
        "aussi une limite serveur configurable, fixee par defaut a 120 secondes."
    )
    add_table(
        doc,
        ["Endpoint", "Methode", "Role"],
        [
            ("/api/health", "GET", "Etat du backend."),
            ("/api/auth/login", "POST", "Connexion."),
            ("/api/auth/register", "POST", "Creation d'un utilisateur."),
            ("/api/auth/me", "GET", "Utilisateur courant."),
            ("/api/meta", "GET", "Modes, moteurs et disponibilite runtime."),
            ("/api/dashboard", "GET", "Statistiques issues de l'historique."),
            ("/api/models", "GET", "Etat des moteurs et modeles."),
            ("/api/history", "GET", "Recherche, filtres et pagination."),
            ("/api/history/{key}", "GET/DELETE", "Detail ou mise en corbeille."),
            ("/api/history/{key}/source", "GET", "Document source archive."),
            ("/api/history/{key}/report.pdf", "GET", "Rapport PDF individuel."),
            ("/api/history/export/zip", "GET", "Export groupe de rapports."),
            ("/api/extractions", "POST", "Traitement d'un batch."),
        ],
        [3650, 1200, 4510],
    )
    add_heading(doc, "7.4 Pretraitement des documents", 2)
    add_bullets(
        doc,
        [
            "Validation de l'extension, de la taille maximale de 20 Mo et du contenu non vide.",
            "Correction de l'orientation EXIF.",
            "Detection du contour principal et transformation a quatre points.",
            "Estimation du deskew par lignes de Hough.",
            "Amelioration du contraste et de la nettete.",
            "Redimensionnement raisonnable pour limiter le cout de calcul.",
            "Preparation des PDF et conservation d'une metadata detaillee des etapes.",
        ],
    )
    add_heading(doc, "7.5 Pipeline OCR local specialise", 2)
    add_paragraph(
        doc,
        "Ce pipeline associe Tesseract, OpenCV, expressions regulieres et heuristiques metier. "
        "Il est privilegie pour STEG et certaines analyses medicales. Des zones rapides peuvent "
        "etre lues pour la reference, le montant et la date. Le medical teste plusieurs variantes "
        "de niveaux de gris et de segmentation, puis choisit la meilleure selon un score interne."
    )
    add_heading(doc, "7.6 Pipeline IA local hybride", 2)
    pipeline_shape = doc.add_picture(str(pipeline_path), width=Cm(16.2))
    set_picture_alt(
        pipeline_shape,
        "Pipeline IA local hybride",
        "Huit etapes allant de la validation du document a son archivage, avec Docling, PaddleOCR, Qwen2.5 et validation metier.",
    )
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Figure 2 - Etapes du pipeline IA local")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.runs[0], size=9, color=MUTED, italic=True)
    add_paragraph(
        doc,
        "Pour les PDF, Docling est tente en premier. Si le Markdown est juge exploitable, il est "
        "envoye directement a Qwen. Sinon, PaddleOCR est lance avec une limite courte; un OCR "
        "Tesseract rapide peut encore prendre le relais. Le prompt impose un schema par famille, "
        "interdit l'invention et demande uniquement un JSON valide. Ollama est appele avec "
        "temperature zero, un contexte configurable et une limite de generation."
    )
    add_callout(
        doc,
        "Detail important",
        "Le mode nomme « local » ne signifie pas que Qwen est toujours utilise. Pour eviter les "
        "timeouts, les images STEG et certaines images medicales suivent des branches OCR rapides. "
        "Le champ local_pipeline indique la source reelle et, dans plusieurs fallbacks, qwen_used=false.",
        fill=LIGHT_GOLD,
        accent=GOLD,
    )
    add_heading(doc, "7.7 Pipeline Gemini", 2)
    add_paragraph(
        doc,
        "Gemini recoit l'image et un prompt specialise. Les pipelines normalisent ensuite la "
        "reponse selon les schemas Pydantic. La cle peut provenir du fichier .env ou etre fournie "
        "pour la session. Ce moteur offre une bonne comprehension visuelle, mais depend du reseau, "
        "du quota, de la politique de confidentialite et du cout du fournisseur."
    )
    add_heading(doc, "7.8 Validation et score qualite", 2)
    add_paragraph(
        doc,
        "Le backend verifie les champs obligatoires, la presence de lignes et la forme des "
        "montants. Il ajoute des warnings lorsque des informations sont absentes ou ambigues. "
        "Un score de completude pondere est calcule differemment selon la famille. Par exemple, "
        "la reference et le montant STEG ont un poids eleve; le medical valorise les identifiants, "
        "les dates et les lignes d'analyse; le ticket valorise le magasin, la date, le total, "
        "l'adresse et le numero. Ce score sert au tri et a l'interface, pas a mesurer une accuracy."
    )
    add_heading(doc, "7.9 Authentification et securite", 2)
    add_paragraph(
        doc,
        "L'authentification est optionnelle et desactivee par defaut pour la demonstration locale. "
        "Lorsqu'elle est active, les mots de passe sont hashes par PBKDF2-SHA256 avec 210 000 "
        "iterations et un sel aleatoire. Le backend emet un token signe en HMAC-SHA256 contenant "
        "le sujet, la date d'emission et l'expiration. Les comptes crees sont stockes dans une base "
        "SQLite distincte. Les secrets doivent rester dans .env."
    )
    add_heading(doc, "7.10 Rapports et historique", 2)
    add_paragraph(
        doc,
        "ReportLab genere un rapport PDF adapte au type de document. Le frontend peut telecharger "
        "un PDF individuel ou demander un ZIP contenant plusieurs rapports. La suppression "
        "deplace le JSON et la source vers une corbeille avant de retirer la ligne active de la "
        "base, ce qui evite une perte immediate."
    )

    add_chapter(doc, 8, "Preparation des donnees et fine-tuning")
    add_heading(doc, "8.1 Objectif", 2)
    add_paragraph(
        doc,
        "Le fine-tuning n'est pas execute pendant chaque extraction. Il constitue une phase "
        "hors ligne destinee a specialiser Qwen2.5. Une fois un adaptateur LoRA entraine et "
        "integre dans Ollama, l'application pourra remplacer le nom du modele sans modifier "
        "l'architecture generale."
    )
    dataset_shape = doc.add_picture(str(dataset_path), width=Cm(15.8))
    set_picture_alt(
        dataset_shape,
        "Repartition du dataset",
        "Graphique des 839 exemples d'entrainement, 178 de validation et 182 de test.",
    )
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Figure 3 - Splits du dataset prepare")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.runs[0], size=9, color=MUTED, italic=True)
    add_heading(doc, "8.2 Composition des donnees", 2)
    add_table(
        doc,
        ["Split", "Total", "Medical", "Tickets"],
        [
            ("Train", "839", "158", "681"),
            ("Validation", "178", "33", "145"),
            ("Test", "182", "35", "147"),
            ("Total", "1 199", "226", "973"),
        ],
        [2400, 1800, 2300, 2860],
        first_col_bold=True,
    )
    add_heading(doc, "8.3 Format d'apprentissage", 2)
    add_paragraph(
        doc,
        "Chaque exemple contient une instruction, un texte d'entree, une sortie JSON de verite "
        "terrain et des metadata. Le split est stratifie par famille avec une graine fixe de 42, "
        "ce qui permet de reproduire la repartition 70 % / 15 % / 15 %."
    )
    add_code_block(
        doc,
        '{"instruction": "Extrais les informations importantes...",\n'
        ' "input": "[METHOD_USED=...]\\n\\n<TEXTE DU DOCUMENT>",\n'
        ' "output": {"document_type": "receipt", "...": "..."},\n'
        ' "metadata": {"id": "...", "document_family": "receipt"}}',
    )
    add_heading(doc, "8.4 Chaine de preparation", 2)
    add_numbered(
        doc,
        [
            "Collecter ou generer les sources et leurs annotations.",
            "Produire un texte d'entree OCR ou structure.",
            "Normaliser la verite terrain vers le schema DocuAI.",
            "Eliminer les exemples invalides ou trop ambigus.",
            "Construire les splits sans melanger train et test.",
            "Exporter le format interne ou le format conversationnel LLaMA-Factory.",
            "Lancer un SFT LoRA sur une machine disposant des ressources necessaires.",
            "Evaluer le modele adapte sur le split test gele avant integration.",
        ],
    )
    add_heading(doc, "8.5 Etat actuel", 2)
    add_callout(
        doc,
        "Etat reel au 15 juin 2026",
        "Le dataset et les scripts de preparation, d'export et d'entrainement sont presents. "
        "Qwen2.5:7b-instruct est le modele local de base. Aucun modele fine-tune ne doit encore "
        "etre presente comme integre ou valide dans l'application.",
        fill=LIGHT_GOLD,
        accent=GOLD,
    )

    add_chapter(doc, 9, "Tests, evaluation et resultats")
    add_heading(doc, "9.1 Strategie de test", 2)
    add_bullets(
        doc,
        [
            "Tests unitaires pytest pour les indices de routage et les fallbacks.",
            "Smoke tests multi-familles pour verifier la detection.",
            "QA d'integration par cas avec attentes sur les champs et les erreurs.",
            "Script quantitatif pour comparer les JSON predits a la verite terrain.",
            "Inspection manuelle des documents difficiles et des warnings.",
        ],
    )
    add_heading(doc, "9.2 Resultats effectivement verifies", 2)
    add_table(
        doc,
        ["Verification", "Date", "Resultat", "Interpretation"],
        [
            ("Detection STEG, medical, ticket, fournisseur", "09/06/2026", "4/4 cas corrects", "Smoke test controle, pas benchmark global."),
            ("QA fonctionnelle recente", "09/06/2026", "7/7 cas passes", "Verifie flux et fallbacks; ne prouve pas la precision globale."),
            ("Tests pytest du routeur", "15/06/2026", "2/2 passes", "Nom d'origine conserve et fallback STEG verifie."),
            ("Split quantitatif", "09/06/2026", "182 documents", "35 medical + 147 tickets."),
            ("Predictions comparables", "15/06/2026", "Aucune disponible", "Precision, Recall et F1 globaux non calculables."),
        ],
        [2900, 1500, 1800, 3160],
    )
    add_heading(doc, "9.3 Metriques retenues", 2)
    add_table(
        doc,
        ["Metrique", "Formule / definition", "Usage"],
        [
            ("Detection Accuracy", "Documents bien classes / documents testes", "Qualite du routeur."),
            ("Valid JSON Rate", "JSON lisibles / sorties produites", "Exploitabilite applicative."),
            ("Field Accuracy", "Champs attendus corrects / champs attendus", "Qualite metier intuitive."),
            ("Precision", "TP / (TP + FP)", "Part des valeurs produites qui sont correctes."),
            ("Recall", "TP / (TP + FN)", "Part des valeurs attendues retrouvees."),
            ("F1-score", "2 x Precision x Recall / (Precision + Recall)", "Equilibre precision/rappel."),
            ("CER / WER", "Distance d'edition / taille de reference", "Qualite du texte OCR."),
            ("Latence", "Temps moyen par document", "Experience utilisateur et cout machine."),
        ],
        [1900, 3700, 3760],
        first_col_bold=True,
    )
    add_heading(doc, "9.4 Analyse des echecs observes", 2)
    add_paragraph(
        doc,
        "Les journaux de QA montrent plusieurs modes d'echec utiles pour l'amelioration. Une "
        "facture STEG a produit un montant OCR incorrect dans un run anterieur, alors que la "
        "branche locale rapide a retrouve la valeur attendue. Une facture fournisseur locale a "
        "echoue lorsque Ollama n'etait pas joignable. Les tickets et factures fournisseurs traites "
        "par fallback OCR peuvent produire un JSON exploitable mais avec un score faible et des "
        "champs a verifier. Ces cas confirment l'interet des warnings, de l'archivage des erreurs "
        "et de la comparaison par famille."
    )
    add_heading(doc, "9.5 Precautions pour le benchmark", 2)
    add_bullets(
        doc,
        [
            "Utiliser exactement les memes documents pour chaque pipeline.",
            "Geler le split test et ne pas regler les prompts sur ce split.",
            "Sauvegarder les predictions sans correction humaine.",
            "Ne pas melanger les exemples synthetiques avec un benchmark clinique reel.",
            "Ne pas compter le score de completude de l'interface comme une accuracy.",
            "Le raccourci receipt_test renvoie l'annotation SROIE lorsque le nom correspond au dataset; il doit etre exclu de toute mesure de modele.",
        ],
    )
    add_heading(doc, "9.6 Decision technique par famille", 2)
    add_table(
        doc,
        ["Type", "Pipeline recommande", "Raison"],
        [
            ("Facture STEG", "OCR specialise; Gemini si image tres difficile", "Structure assez stable, rapidite et explicabilite."),
            ("Analyse medicale", "OCR structure + validation; IA en secours", "Champs critiques et besoin de controle."),
            ("Ticket de caisse", "Gemini ou IA locale hybride", "Mise en page et contenu tres variables."),
            ("Facture fournisseur", "Gemini ou Docling + Qwen", "Semantique, multilingue, lignes et taxes."),
            ("PDF propre", "Docling + Qwen local", "Bonne conservation de la structure."),
        ],
        [2100, 3600, 3660],
        first_col_bold=True,
    )

    add_chapter(doc, 10, "Limites, securite et perspectives")
    add_heading(doc, "10.1 Limites fonctionnelles et scientifiques", 2)
    add_bullets(
        doc,
        [
            "Les metriques F1 ne sont pas encore produites pour les pipelines sur le split test.",
            "La verite terrain couvre principalement medical et tickets; STEG et fournisseur necessitent plus d'annotations.",
            "Les donnees medicales synthetiques ne remplacent pas une validation clinique anonymisee.",
            "Qwen2.5 local n'est pas encore fine-tune ni valide face a Gemini sur un benchmark commun.",
            "La qualite arabe reste dependante des langues OCR installees et de la resolution.",
            "Le batch est traite sequentiellement et aucune file de taches persistante n'est encore utilisee.",
            "Le stockage simultane en BLOB et sur disque augmente l'espace consomme.",
        ],
    )
    add_heading(doc, "10.2 Limites de l'interface actuelle", 2)
    add_paragraph(
        doc,
        "La page Dashboard contient encore des constantes de presentation et une mise a l'echelle "
        "statique pour certaines cartes et graphiques. Les statistiques backend existent, mais "
        "ces valeurs de demonstration doivent etre remplacees par le payload reel avant une "
        "livraison de production ou une presentation comme resultat mesure."
    )
    add_heading(doc, "10.3 Durcissement securite", 2)
    add_bullets(
        doc,
        [
            "Activer DOCUAI_AUTH_ENABLED et fournir un secret de token long et aleatoire.",
            "Utiliser uniquement le hash PBKDF2, jamais un mot de passe en clair en production.",
            "Deplacer le token frontend vers un cookie HttpOnly si l'application devient publique.",
            "Limiter CORS aux domaines deployes et activer HTTPS.",
            "Ajouter une autorisation par role et une politique de conservation des documents.",
            "Chiffrer ou externaliser les documents sensibles.",
            "Auditer les dependances et limiter les informations techniques exposees dans les erreurs.",
        ],
    )
    add_heading(doc, "10.4 Perspectives", 2)
    add_numbered(
        doc,
        [
            "Generer les predictions des trois pipelines sur le split test et calculer les metriques.",
            "Annoter un benchmark STEG et fournisseur plus large.",
            "Entrainer et evaluer un adaptateur LoRA Qwen2.5.",
            "Ajouter Celery/RQ ou une file de taches pour les traitements longs.",
            "Passer a PostgreSQL et a un stockage objet lorsque le volume augmente.",
            "Ajouter une validation humaine et une boucle de correction pour enrichir le dataset.",
            "Mettre en place Docker, CI/CD, tests end-to-end et monitoring.",
            "Ajouter CER/WER et latence par etape dans le tableau d'evaluation.",
            "Remplacer toutes les valeurs de demonstration du dashboard par des donnees reelles.",
        ],
    )

    add_chapter(doc, 11, "Conclusion generale")
    add_paragraph(
        doc,
        "DocuAI constitue une plateforme fonctionnelle et modulaire d'extraction documentaire. "
        "Le projet depasse un simple script OCR : il couvre l'import, la detection, le "
        "pretraitement, plusieurs moteurs, la normalisation, la validation, l'historisation, "
        "l'authentification, l'interface et l'evaluation. L'association Next.js/React et FastAPI "
        "offre une architecture claire, tandis que les services Python permettent de reutiliser "
        "les connaissances metier developpees pour chaque famille."
    )
    add_paragraph(
        doc,
        "Le choix le plus important est l'approche hybride. L'OCR specialise reste pertinent "
        "lorsque la structure est stable; Docling, PaddleOCR et Qwen2.5 offrent une solution locale "
        "pour les documents variables; Gemini fournit une reference cloud. Le systeme enregistre "
        "la methode reelle et les fallbacks afin de ne pas masquer la complexite du traitement."
    )
    add_paragraph(
        doc,
        "Enfin, le projet adopte une communication experimentale prudente. Les resultats "
        "fonctionnels disponibles sont documentes, mais aucun score global n'est invente. La "
        "prochaine etape decisive est la production de predictions comparables sur le split test, "
        "puis l'entrainement et l'evaluation d'un Qwen specialise. Cette base rend le projet "
        "defendable comme PFA et extensible vers une solution plus industrielle."
    )

    add_chapter(doc, 12, "Bibliographie et annexes")
    add_heading(doc, "12.1 Bibliographie et documentation technique", 2)
    references = [
        "React Documentation - https://react.dev/",
        "Next.js Documentation - https://nextjs.org/docs",
        "FastAPI Documentation - https://fastapi.tiangolo.com/",
        "Pydantic Documentation - https://docs.pydantic.dev/",
        "OpenCV Documentation - https://docs.opencv.org/",
        "Tesseract OCR - https://github.com/tesseract-ocr/tesseract",
        "Docling - https://github.com/docling-project/docling",
        "PaddleOCR - https://github.com/PaddlePaddle/PaddleOCR",
        "Ollama Documentation - https://ollama.com/",
        "Qwen Documentation and model cards - https://qwenlm.github.io/",
        "Google Gemini API Documentation - https://ai.google.dev/",
        "SQLite Documentation - https://www.sqlite.org/docs.html",
        "CRISP-DM process model, IBM/SPSS reference materials.",
    ]
    add_bullets(doc, references)

    add_heading(doc, "Annexe A - Commandes de lancement", 2)
    add_code_block(
        doc,
        "python -m pip install -r requirements.txt\n"
        "ollama pull qwen2.5:7b-instruct\n"
        "powershell -ExecutionPolicy Bypass -File .\\run_api.ps1\n"
        "powershell -ExecutionPolicy Bypass -File .\\run_frontend.ps1\n\n"
        "Backend : http://127.0.0.1:8000/api/health\n"
        "Frontend: http://localhost:3000",
    )
    add_heading(doc, "Annexe B - Variables de configuration principales", 2)
    add_table(
        doc,
        ["Variable", "Role", "Valeur typique"],
        [
            ("GEMINI_API_KEY", "Cle API cloud.", "Secret"),
            ("GEMINI_MODEL", "Modele Gemini.", "gemini-2.5-flash"),
            ("OLLAMA_HOST", "Serveur Ollama.", "http://127.0.0.1:11434"),
            ("OLLAMA_MODEL", "Modele local.", "qwen2.5:7b-instruct"),
            ("DOCUAI_EXTRACTION_TIMEOUT_SECONDS", "Timeout document.", "120"),
            ("OLLAMA_GENERATE_TIMEOUT_SECONDS", "Timeout generation.", "90"),
            ("LOCAL_PIPELINE_CONTENT_CHARS", "Taille du contenu transmis.", "16000"),
            ("DOCUAI_AUTH_ENABLED", "Activation auth.", "false en demo"),
            ("DOCUAI_AUTH_TOKEN_TTL_MINUTES", "Duree du token.", "480"),
            ("EXTRACTION_HISTORY_DB_PATH", "Base d'historique.", "Data/history/extractions.db"),
        ],
        [3300, 3660, 2400],
        first_col_bold=True,
    )
    add_heading(doc, "Annexe C - Arborescence simplifiee", 2)
    add_code_block(
        doc,
        "pfaEXTRACT/\n"
        "|-- backend/              API FastAPI, auth et orchestration\n"
        "|-- frontend/             Next.js, React, TypeScript, Tailwind\n"
        "|-- src/\n"
        "|   |-- extraction/       Extracteurs specialises\n"
        "|   |-- services/         Routage, pretraitement, IA, historique, PDF\n"
        "|   |-- models/           Schemas Pydantic\n"
        "|   `-- web/              Ancienne interface Streamlit\n"
        "|-- pipelines/            Pipelines Gemini par famille\n"
        "|-- scripts/              Evaluation, QA et fine-tuning\n"
        "|-- Data/                 Documents, historique et dataset\n"
        "|-- outputs/              Resultats d'evaluation et QA\n"
        "|-- tests/                Tests pytest\n"
        "`-- docs/                 Documentation et rapports",
    )
    add_heading(doc, "Annexe D - Scenario de demonstration", 2)
    add_numbered(
        doc,
        [
            "Lancer Ollama, le backend et le frontend.",
            "Ouvrir la page Extractions et importer une image ou un PDF.",
            "Choisir Auto et le pipeline local, ou selectionner OCR/Gemini pour comparaison.",
            "Lancer l'extraction et observer la redirection vers Documents.",
            "Comparer l'apercu source, les champs, les warnings et le JSON technique.",
            "Telecharger le rapport PDF, puis consulter l'historique et le dashboard.",
            "Ouvrir Evaluation IA pour expliquer le dataset, les metriques et les limites.",
        ],
    )
    add_heading(doc, "Annexe E - Points a personnaliser avant depot", 2)
    add_bullets(
        doc,
        [
            "Completer le nom de l'etudiant, l'encadrant et l'etablissement sur la couverture.",
            "Ajouter le logo officiel de l'etablissement si necessaire.",
            "Inserer des captures d'ecran finales de l'application.",
            "Remplacer les donnees de presentation du dashboard par les statistiques reelles.",
            "Ajouter les resultats F1 uniquement apres execution du benchmark complet.",
            "Relire la terminologie PFA/PFE selon les regles de l'etablissement.",
        ],
    )

    core = doc.core_properties
    core.title = "Rapport PFA detaille - DocuAI / pfaEXTRACT"
    core.subject = "Plateforme intelligente d'extraction d'informations documentaires"
    core.author = "[Nom et prenom]"
    core.keywords = "OCR, DocuAI, Next.js, FastAPI, Qwen, Gemini, PFA"
    core.comments = "Rapport genere a partir du depot pfaEXTRACT le 15 juin 2026."

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_report()
    print(output)
