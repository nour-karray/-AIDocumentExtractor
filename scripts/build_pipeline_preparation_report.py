from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "RAPPORT_PIPELINE_APPLICATION_PREPARATION_DONNEES.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(17, 26, 59)
MUTED = RGBColor(96, 107, 137)
BORDER = "DADCE0"
HEADER_FILL = "F2F4F7"
LIGHT_FILL = "F8FAFD"
CALLOUT_FILL = "EEF4FF"


def set_run_font(run, *, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def add_para(doc, text="", *, style=None, bold=False, italic=False, color=None, size=None, after=6):
    p = doc.add_paragraph(style=style)
    set_paragraph_spacing(p, after=after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    set_paragraph_spacing(
        p,
        before=16 if level == 1 else 12 if level == 2 else 8,
        after=8 if level == 1 else 6 if level == 2 else 4,
    )
    run = p.add_run(text)
    set_run_font(
        run,
        size=16 if level == 1 else 13 if level == 2 else 12,
        color=BLUE if level in (1, 2) else DARK_BLUE,
        bold=True,
    )
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=4, line=1.10)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    run = p.add_run(text)
    set_run_font(run, size=10.7, color=INK)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4, line=1.10)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    run = p.add_run(text)
    set_run_font(run, size=10.7, color=INK)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=2)
    r = p.add_run(title)
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0)
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    add_para(doc, "", after=4)


def add_simple_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], HEADER_FILL)
        p = header_cells[idx].paragraphs[0]
        set_paragraph_spacing(p, after=0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, color=INK, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.08)
            r = p.add_run(value)
            set_run_font(r, size=9.3, color=INK)
    return table


def paragraph_bottom_border(paragraph, color="2E74B5", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, after=0)
    run = p.add_run("Rapport pipeline application | Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True


def build_report():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    kicker = doc.add_paragraph()
    set_paragraph_spacing(kicker, after=6)
    r = kicker.add_run("Rapport technique")
    set_run_font(r, size=11, color=MUTED, bold=True)

    title = doc.add_paragraph()
    set_paragraph_spacing(title, after=4)
    r = title.add_run("Pipeline de l'application pour la préparation des données")
    set_run_font(r, size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, after=14)
    r = subtitle.add_run("Projet DocuAI - Extraction intelligente de documents")
    set_run_font(r, size=13, color=MUTED)

    meta_rows = [
        ["Chapitre concerné", "Compréhension et préparation des données"],
        ["Portée", "Importation, prétraitement, OCR, extraction, JSON, validation et stockage"],
        ["Documents traités", "Factures STEG, analyses médicales, tickets de caisse, factures fournisseurs"],
        ["Pipeline principal", "Docling + PaddleOCR + Qwen local via Ollama"],
        ["Alternative", "API Gemini lorsque cette méthode est sélectionnée"],
    ]
    add_simple_table(doc, ["Élément", "Description"], meta_rows, [2600, 6760])
    rule = doc.add_paragraph()
    paragraph_bottom_border(rule)

    add_heading(doc, "1. Objectif du rapport", 1)
    add_para(
        doc,
        "Ce rapport décrit le pipeline de l'application lié au chapitre de compréhension et de préparation des données. "
        "Il explique comment un document brut est contrôlé, prétraité, analysé par OCR ou par intelligence artificielle, "
        "puis transformé en données JSON exploitables par la plateforme.",
    )
    add_callout(
        doc,
        "Idée centrale",
        "Le rôle de ce pipeline est de transformer des documents hétérogènes et parfois bruités en informations fiables, "
        "normalisées et directement utilisables dans l'interface DocuAI.",
    )

    add_heading(doc, "2. Données et documents concernés", 1)
    add_para(
        doc,
        "La plateforme traite plusieurs familles de documents. Chaque famille possède une structure propre, des champs "
        "importants à extraire et des difficultés spécifiques liées à la qualité de l'image ou à la variabilité du format.",
    )
    doc_rows = [
        [
            "Facture STEG",
            "Référence client, numéro compteur, date facture, montant à payer",
            "Structure stable mais qualité d'image variable",
        ],
        [
            "Analyse médicale",
            "Laboratoire, patient, dossier, dates, médecin, résultats des examens",
            "Tableaux complexes, unités, valeurs numériques et formats variables",
        ],
        [
            "Ticket de caisse",
            "Magasin, date, numéro ticket, articles, total, paiement",
            "Petit texte, papier thermique, bruit et faible contraste",
        ],
        [
            "Facture fournisseur",
            "Fournisseur, client, numéro facture, articles, taxes, total TTC",
            "Mise en page très variable selon l'émetteur",
        ],
    ]
    add_simple_table(doc, ["Type", "Champs principaux", "Difficultés"], doc_rows, [1900, 4100, 3360])

    add_heading(doc, "3. Pipeline global de l'application", 1)
    add_para(
        doc,
        "Le pipeline global suit une chaîne de traitement progressive. Chaque étape prépare la suivante afin de réduire "
        "les erreurs et de produire un résultat structuré.",
    )
    steps = [
        ["1", "Importation", "L'utilisateur dépose un PDF ou une image depuis l'interface web.", "Fichier brut"],
        ["2", "Validation", "Contrôle du format, de la taille et de la lisibilité du document.", "Fichier accepté ou erreur"],
        ["3", "Prétraitement", "Correction de l'orientation, redimensionnement, perspective, inclinaison et contraste.", "Image ou PDF préparé"],
        ["4", "Détection du type", "Identification automatique du document à partir du nom, des mots-clés et du texte OCR.", "Type documentaire"],
        ["5", "Extraction", "Docling, PaddleOCR, Tesseract, Qwen local ou Gemini selon la méthode choisie.", "Texte et champs extraits"],
        ["6", "Structuration", "Conversion des champs extraits en JSON selon un schéma métier.", "Objet JSON"],
        ["7", "Validation", "Contrôle des champs obligatoires et génération d'avertissements.", "Résultat qualifié"],
        ["8", "Stockage", "Archivage du fichier source, du JSON, du statut, du score et de la méthode.", "Historique exploitable"],
    ]
    add_simple_table(doc, ["N°", "Étape", "Traitement", "Sortie"], steps, [520, 1600, 5200, 2040])

    add_heading(doc, "4. Prétraitement des documents avec OpenCV", 1)
    add_para(
        doc,
        "Le prétraitement vise à améliorer la lisibilité du document avant l'OCR. Dans l'application, les images sont "
        "normalisées puis renforcées afin d'augmenter la qualité du texte détecté.",
    )
    preprocessing = [
        ["Validation du format", "Vérifie les extensions autorisées : PDF, JPG, PNG, TIFF, WEBP et BMP."],
        ["Correction d'orientation", "Corrige l'orientation EXIF et teste les rotations possibles lorsque c'est nécessaire."],
        ["Redimensionnement", "Réduit ou agrandit l'image pour obtenir une taille exploitable par l'OCR."],
        ["Correction de perspective", "Détecte le contour du document et redresse la zone utile."],
        ["Correction d'inclinaison", "Utilise les lignes du document pour réduire l'effet de document penché."],
        ["Amélioration du contraste", "Applique CLAHE et renforcement de netteté pour mieux séparer texte et arrière-plan."],
    ]
    add_simple_table(doc, ["Opération", "Rôle dans le pipeline"], preprocessing, [2800, 6560])

    add_heading(doc, "5. Pipeline local : Docling, PaddleOCR et Qwen", 1)
    add_para(
        doc,
        "Le pipeline local est la méthode principale. Il permet de traiter les documents sans envoyer les données vers "
        "un service externe. Il combine extraction documentaire, OCR et modèle de langage local.",
    )
    add_number(doc, "Docling tente de convertir le document en texte ou en Markdown structuré.")
    add_number(doc, "L'application évalue la qualité du contenu extrait : nombre de caractères, mots, lignes, chiffres et tableaux.")
    add_number(doc, "Si Docling est insuffisant, le pipeline bascule vers PaddleOCR.")
    add_number(doc, "Si PaddleOCR échoue ou dépasse le délai, Tesseract OCR est utilisé comme secours rapide.")
    add_number(doc, "Le texte obtenu est envoyé à Qwen local via Ollama avec un schéma JSON attendu.")
    add_number(doc, "Qwen retourne un objet JSON contenant les champs métier détectés.")

    local_rows = [
        ["Docling", "Extraction de contenu documentaire structuré, surtout utile pour les PDF."],
        ["PaddleOCR", "OCR robuste utilisé comme fallback lorsque Docling ne fournit pas assez de texte."],
        ["Tesseract OCR", "Secours local rapide pour obtenir du texte minimal exploitable."],
        ["Qwen local", "Transformation du texte en JSON structuré selon le type de document."],
        ["Ollama", "Serveur local qui exécute le modèle Qwen sur la machine."],
    ]
    add_simple_table(doc, ["Composant", "Rôle"], local_rows, [2100, 7260])

    add_heading(doc, "6. Pipeline Gemini", 1)
    add_para(
        doc,
        "L'application conserve une méthode basée sur l'API Gemini. Cette méthode est utilisée lorsque l'utilisateur la "
        "sélectionne explicitement, notamment pour comparer les résultats avec le pipeline local ou traiter certains "
        "documents complexes.",
    )
    add_callout(
        doc,
        "Différence principale",
        "Le pipeline local garde les données dans l'environnement de l'utilisateur, alors que Gemini s'appuie sur une API "
        "externe. Les résultats Gemini sont ensuite normalisés pour rester compatibles avec les mêmes pages Documents, "
        "Historiques et Tableau de bord.",
    )

    add_heading(doc, "7. Normalisation et résultat JSON", 1)
    add_para(
        doc,
        "Après l'extraction, les champs sont convertis dans une structure JSON standardisée. Cette normalisation rend "
        "l'affichage, le stockage, la recherche et l'export PDF cohérents pour tous les types de documents.",
    )
    json_table = doc.add_table(rows=1, cols=1)
    set_table_geometry(json_table, [9360])
    set_table_borders(json_table)
    cell = json_table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.0)
    code = (
        '{\n'
        '  "document_type": "steg_invoice",\n'
        '  "reference": "123456789",\n'
        '  "numero_compteur": "456789",\n'
        '  "date_facture": "2025-05-12",\n'
        '  "montant_a_payer": "185.750",\n'
        '  "warnings": []\n'
        '}'
    )
    r = p.add_run(code)
    set_run_font(r, name="Consolas", size=9.3, color=INK)

    add_heading(doc, "8. Validation métier et score qualité", 1)
    add_para(
        doc,
        "Avant l'enregistrement, l'application vérifie que les champs importants sont présents. Cette étape évite de "
        "considérer comme satisfaisant un résultat incomplet.",
    )
    validations = [
        ["Facture STEG", "Référence, numéro compteur, date facture, montant à payer"],
        ["Analyse médicale", "Patient et résultats des examens"],
        ["Ticket de caisse", "Magasin et total"],
        ["Facture fournisseur", "Numéro facture et total"],
    ]
    add_simple_table(doc, ["Type de document", "Champs obligatoires contrôlés"], validations, [2500, 6860])
    add_para(
        doc,
        "Un score qualité est calculé à partir des champs détectés. Dans l'interface, ce score permet de classer le "
        "résultat comme fiable, correct ou à vérifier.",
    )

    add_heading(doc, "9. Stockage et exploitation dans l'application", 1)
    add_para(
        doc,
        "Une fois le traitement terminé, le fichier source et le JSON sont archivés dans l'historique. Les informations "
        "sont ensuite disponibles dans les pages Documents, Historiques et Tableau de bord.",
    )
    storage_rows = [
        ["Fichier source", "Permet de consulter le document original après extraction."],
        ["Résultat JSON", "Contient les champs extraits et les métadonnées techniques."],
        ["Méthode utilisée", "Indique si le document vient de l'OCR local, du pipeline local ou de Gemini."],
        ["Statut", "Indique si le traitement est en succès ou en erreur."],
        ["Rapport PDF", "Permet l'export du résultat sous forme de rapport lisible."],
    ]
    add_simple_table(doc, ["Élément stocké", "Utilité"], storage_rows, [2500, 6860])

    add_heading(doc, "10. Synthèse", 1)
    add_para(
        doc,
        "Le pipeline de l'application associe le prétraitement d'image, l'OCR et l'intelligence artificielle afin de "
        "transformer des documents hétérogènes en données structurées. Cette approche hybride améliore la robustesse "
        "du système face aux documents flous, inclinés, bruités ou présentant des mises en page différentes.",
    )
    add_para(
        doc,
        "Ce pipeline constitue la base technique nécessaire aux étapes suivantes du projet : choix des modèles, "
        "évaluation des performances, comparaison entre OCR local, Docling, PaddleOCR, Qwen local et Gemini, puis "
        "exploitation des résultats dans l'interface web.",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_report())
